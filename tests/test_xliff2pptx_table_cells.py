"""Issue #58: positional table-cell backfill for xliff2pptx.

OPP emits one trans-unit per PPTX table cell with resname
``table_{t}_r{r}_c{c}``, where ``t`` accumulates across slides in presentation
order.  ORF's old text-matching pass keys a source-text dict, so a value
repeated across cells collapses to ONE target and no ordering guarantee exists
for the ``t`` accumulator.  These tests pin the positional write-back, the
``sldIdLst`` ordering, and the fallbacks.
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from unittest.mock import patch

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from orf.channels.xliff2pptx import XLIFF2PPTXConverter

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

A = f"{{{A_NS}}}"
P = f"{{{P_NS}}}"
R_ID = f"{{{R_NS}}}id"


# ── XLIFF fixtures (1.2 — the version that carries resname) ────────────────


def _xliff(units: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">\n'
        '  <file original="deck" source-language="en" target-language="zh">\n'
        "    <body>\n" + units + "\n    </body>\n  </file>\n</xliff>\n"
    )


def _unit(uid: str, resname: str, source: str, target: str) -> str:
    return (
        f'      <trans-unit xml:space="preserve" id="{uid}" resname="{resname}">\n'
        f'        <source xml:space="preserve">{source}</source>\n'
        f'        <target xml:space="preserve">{target}</target>\n'
        "      </trans-unit>"
    )


def _plain_unit(uid: str, source: str, target: str) -> str:
    return (
        f'      <trans-unit xml:space="preserve" id="{uid}">\n'
        f'        <source xml:space="preserve">{source}</source>\n'
        f'        <target xml:space="preserve">{target}</target>\n'
        "      </trans-unit>"
    )


# ── PPTX fixtures ─────────────────────────────────────────────────────────


def _build_deck(path: Path, slides: list[dict]) -> Path:
    """Build a real .pptx.  Each slide spec: {"tables": [(rows, cols, cells)],
    "textboxes": [str]}"""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for spec in slides:
        slide = prs.slides.add_slide(blank)
        for rows, cols, cells in spec.get("tables", []):
            graphic_frame = slide.shapes.add_table(
                rows, cols, Inches(1), Inches(1), Inches(4), Inches(2)
            )
            table = graphic_frame.table
            for (row, col), text in cells.items():
                table.cell(row, col).text = text
        for text in spec.get("textboxes", []):
            box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
            box.text_frame.text = text
    prs.save(str(path))
    return path


def _minimal_skeleton(path: Path) -> Path:
    """Minimal PPTX ZIP with no ppt/presentation.xml (legacy fixture shape)."""
    slide_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r>"
        "<a:t>Hello World</a:t>"
        "</a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("ppt/slides/slide1.xml", slide_xml)
    return path


def _reverse_sldid_order(src: Path, dst: Path) -> Path:
    """Swap the first two p:sldId r:id refs so sldIdLst order != filename order."""
    with zipfile.ZipFile(src) as zin:
        data = {name: zin.read(name) for name in zin.namelist()}
    pres = etree.fromstring(data["ppt/presentation.xml"])
    sld_ids = pres.xpath(".//p:sldIdLst/p:sldId", namespaces={"p": P_NS})
    assert len(sld_ids) == 2, "reversal helper expects a 2-slide deck"
    first, second = sld_ids[0].get(R_ID), sld_ids[1].get(R_ID)
    sld_ids[0].set(R_ID, second)
    sld_ids[1].set(R_ID, first)
    data["ppt/presentation.xml"] = etree.tostring(
        pres, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in data.items():
            zout.writestr(name, content)
    return dst


# ── read-back helpers ───────────────────────────────────────────────────────


def _read_slide(path: Path, name: str) -> etree._Element:
    with zipfile.ZipFile(path) as zf:
        return etree.fromstring(zf.read(name))


def _cells(root: etree._Element) -> dict[tuple[int, int, int], str]:
    out: dict[tuple[int, int, int], str] = {}
    for t_idx, tbl in enumerate(root.iter(f"{A}tbl")):
        for r_idx, tr in enumerate(tbl.findall(f"{A}tr")):
            for c_idx, tc in enumerate(tr.findall(f"{A}tc")):
                out[(t_idx, r_idx, c_idx)] = "".join(
                    el.text or "" for el in tc.iter(f"{A}t")
                )
    return out


def _run_texts(root: etree._Element) -> list[str]:
    return [
        "".join(el.text or "" for el in run.iter(f"{A}t"))
        for run in root.iter(f"{A}r")
    ]


def _reopen_tables(path: Path) -> list[list[list[str]]]:
    """Reopen a .pptx and return, per slide (sldIdLst order), each table grid."""
    prs = Presentation(str(path))
    slides: list[list[list[str]]] = []
    for slide in prs.slides:
        grids: list[list[list[str]]] = []
        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                grids.append([
                    [table.cell(r, c).text for c in range(len(table.columns))]
                    for r in range(len(table.rows))
                ])
        slides.append(grids)
    return slides


def _convert(skeleton: Path, tmp_path: Path, xliff_content: str) -> tuple[Path, object]:
    xliff = tmp_path / "translation.xlf"
    xliff.write_text(xliff_content, encoding="utf-8")
    output = tmp_path / "out.pptx"
    result = XLIFF2PPTXConverter().convert(skeleton, xliff, output)
    assert result.success is True, result.errors
    return output, result


# ── tests ───────────────────────────────────────────────────────────────────


class TestParseTableResname:
    def test_table_resname_parsed(self):
        converter = XLIFF2PPTXConverter()
        assert converter._parse_table_resname("table_0_r1_c2") == (0, 1, 2)

    def test_malformed_table_resname_is_none(self):
        converter = XLIFF2PPTXConverter()
        assert converter._parse_table_resname("table_x_r1_c0") is None

    def test_non_table_resname_is_none(self):
        converter = XLIFF2PPTXConverter()
        assert converter._parse_table_resname("para_index_3") is None
        assert converter._parse_table_resname(None) is None


class TestPositionalTableCellBackfill:
    def test_duplicate_cell_text_each_gets_own_target_and_text_pass_skips(
        self, tmp_path: Path
    ):
        # Trap #1: a same-source non-table run must NOT receive a table target.
        skeleton = _build_deck(
            tmp_path / "deck.pptx",
            [{
                "tables": [(2, 2, {(0, 0): "IP67", (1, 0): "IP67",
                                   (0, 1): "Yes", (1, 1): "No"})],
                "textboxes": ["IP67"],
            }],
        )
        xliff = _xliff("\n".join([
            _unit("1", "table_0_r0_c0", "IP67", "ZH-PRIMARY-IP67"),
            _unit("2", "table_0_r1_c0", "IP67", "ZH-SECONDARY-IP67"),
        ]))
        output, _ = _convert(skeleton, tmp_path, xliff)

        grids = _reopen_tables(output)
        assert grids[0][0][0][0] == "ZH-PRIMARY-IP67"
        assert grids[0][0][1][0] == "ZH-SECONDARY-IP67"

        root = _read_slide(output, "ppt/slides/slide1.xml")
        assert _run_texts(root).count("IP67") == 1, (
            "the un-unit'ed text box sharing the table source must be untouched"
        )

    def test_two_slides_tables_map_to_right_slides(self, tmp_path: Path):
        skeleton = _build_deck(
            tmp_path / "deck.pptx",
            [
                {"tables": [(1, 1, {(0, 0): "A1"})]},
                {"tables": [(1, 1, {(0, 0): "B1"})]},
            ],
        )
        xliff = _xliff("\n".join([
            _unit("1", "table_0_r0_c0", "A1", "T-A"),
            _unit("2", "table_1_r0_c0", "B1", "T-B"),
        ]))
        output, _ = _convert(skeleton, tmp_path, xliff)

        assert _cells(_read_slide(output, "ppt/slides/slide1.xml"))[(0, 0, 0)] == "T-A"
        assert _cells(_read_slide(output, "ppt/slides/slide2.xml"))[(0, 0, 0)] == "T-B"

    def test_sldid_order_wins_over_filename_order(self, tmp_path: Path):
        deck = _build_deck(
            tmp_path / "deck.pptx",
            [
                {"tables": [(1, 1, {(0, 0): "AAA"})]},
                {"tables": [(1, 1, {(0, 0): "BBB"})]},
            ],
        )
        reversed_deck = _reverse_sldid_order(deck, tmp_path / "reversed.pptx")
        xliff = _xliff("\n".join([
            _unit("1", "table_0_r0_c0", "AAA", "FIRST"),
            _unit("2", "table_1_r0_c0", "BBB", "SECOND"),
        ]))
        output, _ = _convert(reversed_deck, tmp_path, xliff)

        # sldIdLst now lists slide2 first, so table_0 belongs to slide2.
        assert _cells(_read_slide(output, "ppt/slides/slide2.xml"))[(0, 0, 0)] == "FIRST"
        assert _cells(_read_slide(output, "ppt/slides/slide1.xml"))[(0, 0, 0)] == "SECOND"

    def test_out_of_range_table_warns_without_raising(self, tmp_path: Path):
        skeleton = _build_deck(
            tmp_path / "deck.pptx",
            [{"tables": [(1, 1, {(0, 0): "X"})]}],
        )
        xliff = _xliff(_unit("1", "table_9_r9_c9", "X", "T"))
        with patch("orf.channels.xliff2pptx.logger") as mock_logger:
            output, _ = _convert(skeleton, tmp_path, xliff)

        warned = " ".join(
            str(arg)
            for call in mock_logger.warning.call_args_list
            for arg in call.args
        )
        assert "out of range" in warned
        # Output still opens and the source is untouched (table units never
        # fall back to the text pass).
        assert _reopen_tables(output)[0][0][0][0] == "X"

    def test_multi_paragraph_cell_replaces_all_source(self, tmp_path: Path):
        skeleton = tmp_path / "deck.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(
            1, 1, Inches(1), Inches(1), Inches(4), Inches(2)
        ).table
        frame = table.cell(0, 0).text_frame
        frame.text = "IP"
        frame.add_paragraph().text = "67"
        prs.save(str(skeleton))

        xliff = _xliff(_unit("1", "table_0_r0_c0", "IP\n67", "翻译"))
        output, _ = _convert(skeleton, tmp_path, xliff)

        cell_t = [
            el.text or ""
            for tc in _read_slide(output, "ppt/slides/slide1.xml").iter(f"{A}tc")
            for el in tc.iter(f"{A}t")
        ]
        assert "".join(cell_t) == "翻译"
        assert cell_t[1:] == [""] * (len(cell_t) - 1), "tail runs must be cleared"


class TestTextPathFallback:
    def test_paragraph_only_xliff_still_backfills(self, tmp_path: Path):
        skeleton = _build_deck(
            tmp_path / "deck.pptx", [{"textboxes": ["Hello World"]}]
        )
        xliff = _xliff(_plain_unit("1", "Hello World", "你好 世界"))
        output, _ = _convert(skeleton, tmp_path, xliff)

        root = _read_slide(output, "ppt/slides/slide1.xml")
        assert "你好 世界" in _run_texts(root)

    def test_minimal_fixture_without_presentation_falls_back(self, tmp_path: Path):
        skeleton = _minimal_skeleton(tmp_path / "skeleton.pptx")
        xliff = _xliff(_plain_unit("1", "Hello World", "你好 世界"))
        with patch("orf.channels.xliff2pptx.logger") as mock_logger:
            output, _ = _convert(skeleton, tmp_path, xliff)

        warned = " ".join(
            str(arg)
            for call in mock_logger.warning.call_args_list
            for arg in call.args
        )
        assert "falling back" in warned
        root = _read_slide(output, "ppt/slides/slide1.xml")
        assert "你好 世界" in _run_texts(root)
