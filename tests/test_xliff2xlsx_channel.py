"""XLIFF → XLSX backfill channel tests (ORF half of Omni_Pre_Processor#92).

The skeleton under test is OPP's: the original workbook bytes plus one
archive-root sidecar, ``xliff_map.json``. These tests build that same package
from a *rich* workbook — formula cells, a merged range, two sheets, numeric and
whitespace-only cells, a multi-column row and a row whose edge cells are blank —
because a single-cell fixture proves neither of the two things that can go
wrong in this channel: correlating ``" | "`` segments with cells, and leaving
``translatable: false`` cells alone.

Row sources are computed with OPP's own rule (``" | ".join(cell texts).strip()``)
rather than typed by hand, so a change in that contract surfaces as a failing
test instead of a stale fixture that still agrees with the implementation.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

from orf.channels.xliff2xlsx import XLIFF2XLSXConverter
from orf.converters.base import ConversionResult
from orf.converters.options import ConverterOptions

XLIFF_MAP_ENTRY = "xliff_map.json"

# The map exactly as OPP writes it: ids have gaps because a sheet-header
# paragraph consumes one without recording a row (id 1 = Sheet1 header,
# id 7 = Second header), so rows can only be found by id.
MAP_SPEC: dict[str, list[tuple[int, int, list[tuple[str, bool]]]]] = {
    "Sheet1": [
        (2, 1, [("A1", True), ("B1", True), ("C1", True),
                ("D1", False), ("E1", False)]),
        (3, 2, [("A2", True), ("B2", False), ("C2", False),
                ("D2", False), ("E2", False)]),
        (4, 3, [("A3", False), ("B3", True), ("C3", False),
                ("D3", False), ("E3", False)]),
        (5, 4, [("A4", False), ("B4", False), ("C4", False),
                ("D4", False), ("E4", False)]),
        (6, 5, [("A5", True), ("B5", False), ("C5", False),
                ("D5", False), ("E5", False)]),
    ],
    "Second": [
        (8, 1, [("A1", True)]),
    ],
}

# Translated segments, one per cell of the mapped row, in the same order.
TARGET_SEGMENTS: dict[int, list[str]] = {
    2: ["名称", "数量", "价格", "", "3.5"],
    3: ["小部件", "42", "", "", ""],
    4: ["", "配件", "", "", ""],
    5: ["", "", "", "", ""],
    6: ["合并标题", "", "", "", ""],
    8: ["第二个工作表文本"],
}

# OPP numbers the sheet-header paragraph too, so the raw XLIFF carries units
# (1 and 7) that the map never references. They must be ignored, not written.
HEADER_UNITS = {
    "1": ("=== Sheet: Sheet1 ===", "=== 工作表: Sheet1 ==="),
    "7": ("=== Sheet: Second ===", "=== 工作表: Second ==="),
}


def _row_text(values: list[str]) -> str:
    """OPP's row rule: one " | "-joined segment per cell, then stripped."""
    return " | ".join(values).strip()


def _build_workbook(path: Path) -> None:
    """A workbook whose every hazard for this channel is present."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Name"
    ws["B1"] = "Qty"
    ws["C1"] = "Price"
    ws["E1"] = 3.5
    ws["E1"].number_format = "#,##0.00"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "Widget"
    ws["B2"] = 42
    ws["C2"] = "=B2*2"
    ws["D2"] = "   "
    ws["E2"] = "=1/0"
    ws["B3"] = "Gadget"
    ws.merge_cells("A5:B5")
    ws["A5"] = "Merged Title"
    ws.column_dimensions["A"].width = 30
    wb.create_sheet("Second")["A1"] = "Second sheet text"
    wb.save(path)


def _map_payload() -> dict:
    return {
        "version": 1,
        "sheets": {
            sheet: [
                {
                    "id": unit_id,
                    "row": row,
                    "cells": [
                        {"ref": ref, "translatable": translatable}
                        for ref, translatable in cells
                    ],
                }
                for unit_id, row, cells in rows
            ]
            for sheet, rows in MAP_SPEC.items()
        },
    }


def _make_skeleton(workbook: Path, target: Path) -> Path:
    """Package the workbook verbatim plus the sidecar, as OPP does."""
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zf:
        with zipfile.ZipFile(workbook) as source:
            for info in source.infolist():
                zf.writestr(info, source.read(info))
        zf.writestr(XLIFF_MAP_ENTRY, json.dumps(_map_payload(), ensure_ascii=False))
    return target


def _sheet_values(workbook: Path, sheet: str, row: int, columns: int) -> list[str]:
    """The per-cell texts OPP would join into that row's trans-unit source."""
    ws = openpyxl.load_workbook(workbook)[sheet]
    values = []
    for column in range(1, columns + 1):
        value = ws.cell(row=row, column=column).value
        values.append("" if value is None else str(value))
    return values


def _xliff_text(units: dict[str, tuple[str, str | None]]) -> str:
    body = []
    for unit_id, (source, target) in units.items():
        entry = [f'<trans-unit id="{unit_id}">', f"<source>{source}</source>"]
        if target is not None:
            entry.append(f"<target>{target}</target>")
        entry.append("</trans-unit>")
        body.append("".join(entry))
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2" version="1.2">\n'
        '  <file original="book.xlsx" source-language="en" datatype="plaintext"'
        ' target-language="zh">\n'
        "    <body>\n      " + "\n      ".join(body) + "\n    </body>\n"
        "  </file>\n</xliff>\n"
    )


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    path = tmp_path / "book.xlsx"
    _build_workbook(path)
    return path


@pytest.fixture
def skeleton(workbook: Path, tmp_path: Path) -> Path:
    return _make_skeleton(workbook, tmp_path / "book.skeleton.zip")


@pytest.fixture
def translated_xliff(workbook: Path, tmp_path: Path) -> Path:
    """Every mapped row translated, with OPP's exact row shape."""
    units: dict[str, tuple[str, str | None]] = dict(HEADER_UNITS)
    for sheet, rows in MAP_SPEC.items():
        for unit_id, row, cells in rows:
            source = _row_text(_sheet_values(workbook, sheet, row, len(cells)))
            target = " | ".join(TARGET_SEGMENTS[unit_id])
            units[str(unit_id)] = (source, target)
    path = tmp_path / "book.zh.xlf"
    path.write_text(_xliff_text(units), encoding="utf-8")
    return path


def _convert(
    skeleton: Path, xliff: Path, output: Path
) -> ConversionResult:
    return XLIFF2XLSXConverter().convert(
        skeleton, xliff, output, options=ConverterOptions()
    )


class TestChannelSurface:
    def test_supported_format(self):
        assert XLIFF2XLSXConverter().supported_format == "XLSX"

    def test_validate_input_accepts_skeleton_zip(self, skeleton: Path):
        assert XLIFF2XLSXConverter().validate_input(skeleton) is True

    def test_validate_input_accepts_bare_xlsx(self, workbook: Path):
        assert XLIFF2XLSXConverter().validate_input(workbook) is True

    def test_validate_input_rejects_other_extension(self, tmp_path: Path):
        other = tmp_path / "book.txt"
        other.write_text("x", encoding="utf-8")
        assert XLIFF2XLSXConverter().validate_input(other) is False

    def test_validate_input_rejects_missing_file(self, tmp_path: Path):
        assert XLIFF2XLSXConverter().validate_input(
            tmp_path / "absent.zip"
        ) is False


class TestRoundTripFidelity:
    """What a reader of the output must still find, asserted explicitly.

    The accepted cost of this channel is openpyxl's round trip, which drops
    charts, images, pivot tables, data validation and part of the conditional
    formatting. The assertions below pin what *is* preserved, so the loss
    cannot quietly widen.
    """

    def test_translated_cells_are_written(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        output = tmp_path / "out.xlsx"
        result = _convert(skeleton, translated_xliff, output)

        assert result.success, [e.message for e in result.errors]
        assert output.exists()
        wb = openpyxl.load_workbook(output)
        assert wb["Sheet1"]["A1"].value == "名称"
        assert wb["Sheet1"]["B1"].value == "数量"
        assert wb["Sheet1"]["C1"].value == "价格"
        assert wb["Sheet1"]["A2"].value == "小部件"
        assert wb["Sheet1"]["A5"].value == "合并标题"
        assert wb["Second"]["A1"].value == "第二个工作表文本"
        # A1,B1,C1 + A2 + B3 + A5 on Sheet1, then Second!A1.
        assert result.metadata["cells_written"] == 7

    def test_formula_cells_keep_their_formulas(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        """translatable:false covers formula cells — writing one would drop it."""
        output = tmp_path / "out.xlsx"
        assert _convert(skeleton, translated_xliff, output).success

        wb = openpyxl.load_workbook(output)
        assert wb["Sheet1"]["C2"].value == "=B2*2"
        assert wb["Sheet1"]["E2"].value == "=1/0"

    def test_non_translatable_cells_are_not_clobbered(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        output = tmp_path / "out.xlsx"
        assert _convert(skeleton, translated_xliff, output).success

        wb = openpyxl.load_workbook(output)
        assert wb["Sheet1"]["B2"].value == 42, "numeric cell was overwritten"
        assert wb["Sheet1"]["E1"].value == 3.5, "numeric header was overwritten"
        assert wb["Sheet1"]["D2"].value == "   ", (
            "whitespace-only cell was overwritten"
        )
        assert wb["Sheet1"]["B5"].value is None, (
            "merged-range non-anchor cell was written"
        )
        assert wb["Sheet1"]["A3"].value is None
        assert wb["Sheet1"]["C3"].value is None

    def test_merged_range_survives(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        output = tmp_path / "out.xlsx"
        assert _convert(skeleton, translated_xliff, output).success

        ranges = openpyxl.load_workbook(output)["Sheet1"].merged_cells.ranges
        assert [str(r) for r in ranges] == ["A5:B5"]

    def test_structure_matches_the_source_workbook(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path,
        workbook: Path,
    ):
        output = tmp_path / "out.xlsx"
        assert _convert(skeleton, translated_xliff, output).success

        source = openpyxl.load_workbook(workbook)
        result = openpyxl.load_workbook(output)
        assert result.sheetnames == source.sheetnames
        for name in source.sheetnames:
            assert result[name].dimensions == source[name].dimensions
        assert result["Sheet1"]["A1"].font.bold is True
        assert result["Sheet1"]["E1"].number_format == "#,##0.00"
        assert result["Sheet1"].column_dimensions["A"].width == 30

    def test_output_is_a_valid_xlsx_package(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        output = tmp_path / "out.xlsx"
        assert _convert(skeleton, translated_xliff, output).success

        with zipfile.ZipFile(output) as zf:
            assert zf.testzip() is None
            assert "xl/workbook.xml" in zf.namelist()
            assert XLIFF_MAP_ENTRY not in zf.namelist(), (
                "the sidecar must not leak into the output package"
            )
        assert zipfile.is_zipfile(output)


class TestSegmentToCellCorrelation:
    def test_row_with_blank_edge_cells_lands_in_the_right_cell(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        """OPP strips the joined row, so a blank first cell leaves a bare "|"."""
        row3 = _row_text(
            _sheet_values(_source_of(skeleton), "Sheet1", 3, 5)
        )
        assert row3 == "| Gadget |  |  |", (
            f"fixture no longer reproduces OPP's stripped row: {row3!r}"
        )

        output = tmp_path / "out.xlsx"
        result = _convert(skeleton, translated_xliff, output)
        assert result.success, [e.message for e in result.errors]
        wb = openpyxl.load_workbook(output)
        assert wb["Sheet1"]["B3"].value == "配件"
        assert wb["Sheet1"]["A3"].value is None
        assert wb["Sheet1"]["C3"].value is None

    def test_row_whose_cell_text_contains_the_separator_is_rejected(
        self, tmp_path: Path
    ):
        """A cell holding " | " makes the row text ambiguous — refuse, don't guess."""
        wb = openpyxl.Workbook()
        wb.active.title = "Pipes"
        wb.active["A1"] = "A | B"
        book = tmp_path / "pipes.xlsx"
        wb.save(book)

        payload = {
            "version": 1,
            "sheets": {
                "Pipes": [
                    {"id": 1, "row": 1, "cells": [{"ref": "A1", "translatable": True}]}
                ]
            },
        }
        skel = tmp_path / "pipes.skeleton.zip"
        with zipfile.ZipFile(skel, "w", zipfile.ZIP_DEFLATED) as zf:
            with zipfile.ZipFile(book) as source:
                for info in source.infolist():
                    zf.writestr(info, source.read(info))
            zf.writestr(XLIFF_MAP_ENTRY, json.dumps(payload))

        xlf = tmp_path / "pipes.xlf"
        xlf.write_text(_xliff_text({"1": ("A | B", "甲 | 乙")}), encoding="utf-8")
        output = tmp_path / "out.xlsx"

        result = _convert(skel, xlf, output)
        assert result.success is False
        assert "segment" in result.errors[0].message
        assert not output.exists()

    def test_target_that_drops_the_separator_is_rejected(
        self, skeleton: Path, tmp_path: Path
    ):
        """A translator that merges two segments breaks the positional index."""
        units = dict(HEADER_UNITS)
        for sheet, rows in MAP_SPEC.items():
            for unit_id, row, cells in rows:
                units[str(unit_id)] = (
                    _row_text(_sheet_values(
                        _source_of(skeleton), sheet, row, len(cells)
                    )),
                    None,
                )
        units["2"] = (_row_text(_sheet_values(_source_of(skeleton), "Sheet1", 1, 5)),
                      "名称 数量 价格 3.5")
        xlf = tmp_path / "merged.xlf"
        xlf.write_text(_xliff_text(units), encoding="utf-8")
        output = tmp_path / "out.xlsx"

        result = _convert(skeleton, xlf, output)
        assert result.success is False
        assert "cell structure" in result.errors[0].message
        assert not output.exists()

    def test_inline_elements_do_not_truncate_a_row(
        self, tmp_path: Path
    ):
        """<bx/>/<ex/> children carry no text but must not drop their tails."""
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.active["A1"] = "Name"
        wb.active["B1"] = "Qty"
        book = tmp_path / "book.xlsx"
        wb.save(book)
        payload = {
            "version": 1,
            "sheets": {
                "Sheet1": [
                    {"id": 1, "row": 1, "cells": [
                        {"ref": "A1", "translatable": True},
                        {"ref": "B1", "translatable": True},
                    ]}
                ]
            },
        }
        skel = tmp_path / "book.skeleton.zip"
        with zipfile.ZipFile(skel, "w", zipfile.ZIP_DEFLATED) as zf:
            with zipfile.ZipFile(book) as source:
                for info in source.infolist():
                    zf.writestr(info, source.read(info))
            zf.writestr(XLIFF_MAP_ENTRY, json.dumps(payload))
        xlf = tmp_path / "inline.xlf"
        xlf.write_text(
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2" version="1.2">\n'
            "  <file><body>\n"
            '    <trans-unit id="1"><source>Name | <bx id="b1" type="bold"/>Qty'
            '<ex id="b1"/></source>'
            "<target>名称 | <bx id=\"b1\" type=\"bold\"/>数量<ex id=\"b1\"/>"
            "</target></trans-unit>\n"
            "  </body></file>\n</xliff>\n",
            encoding="utf-8",
        )
        output = tmp_path / "out.xlsx"

        result = _convert(skel, xlf, output)
        assert result.success, [e.message for e in result.errors]
        wb = openpyxl.load_workbook(output)
        assert wb["Sheet1"]["A1"].value == "名称"
        assert wb["Sheet1"]["B1"].value == "数量"


class TestIdLookup:
    def test_second_sheet_is_found_by_id_despite_the_gaps(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        """The Second sheet's unit is id 8 while id 7 is a header gap."""
        output = tmp_path / "out.xlsx"
        result = _convert(skeleton, translated_xliff, output)

        assert result.success, [e.message for e in result.errors]
        assert openpyxl.load_workbook(output)["Second"]["A1"].value == (
            "第二个工作表文本"
        )

    def test_unreferenced_sheet_header_units_are_ignored(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        """ids 1 and 7 carry a sheet header and no cells; writing them is wrong."""
        output = tmp_path / "out.xlsx"
        assert _convert(skeleton, translated_xliff, output).success

        wb = openpyxl.load_workbook(output)
        assert "=== Sheet" not in str(wb["Sheet1"]["A1"].value)
        assert wb["Sheet1"].max_row == 5, "a header unit created a stray row"

    def test_missing_trans_unit_is_reported_without_writing(
        self, skeleton: Path, translated_xliff: Path, tmp_path: Path
    ):
        units = _units_of(translated_xliff)
        del units["6"]
        xlf = tmp_path / "dropped.xlf"
        xlf.write_text(_xliff_text(units), encoding="utf-8")
        output = tmp_path / "out.xlsx"

        result = _convert(skeleton, xlf, output)
        assert result.success is False
        assert "trans-unit with id '6'" in result.errors[0].message
        assert not output.exists()


class TestUntranslatedUnits:
    def test_unit_without_target_keeps_the_original_text(
        self, tmp_path: Path, skeleton: Path, translated_xliff: Path
    ):
        units = _units_of(translated_xliff)
        units["6"] = (units["6"][0], None)
        xlf = tmp_path / "partial.xlf"
        xlf.write_text(_xliff_text(units), encoding="utf-8")
        output = tmp_path / "out.xlsx"

        result = _convert(skeleton, xlf, output)
        assert result.success, [e.message for e in result.errors]
        wb = openpyxl.load_workbook(output)
        assert wb["Sheet1"]["A5"].value == "Merged Title"
        assert wb["Sheet1"]["A1"].value == "名称", "other rows must still be written"
        assert result.metadata["rows_untranslated"] == 1
        assert any(
            "no <target>" in w.message for w in result.warnings
        ), [w.message for w in result.warnings]


class TestSkeletonRejection:
    def test_zip_without_the_sidecar_is_rejected(self, workbook: Path, tmp_path: Path):
        skel = tmp_path / "plain.skeleton.zip"
        with zipfile.ZipFile(skel, "w", zipfile.ZIP_DEFLATED) as zf:
            with zipfile.ZipFile(workbook) as source:
                for info in source.infolist():
                    zf.writestr(info, source.read(info))
        xlf = tmp_path / "t.xlf"
        xlf.write_text(_xliff_text({"1": ("Name", "名称")}), encoding="utf-8")
        output = tmp_path / "out.xlsx"

        result = _convert(skel, xlf, output)
        assert result.success is False
        assert "no xliff_map.json sidecar" in result.errors[0].message
        assert not output.exists()

    def test_bare_workbook_is_rejected_with_a_clear_message(
        self, workbook: Path, translated_xliff: Path, tmp_path: Path
    ):
        output = tmp_path / "out.xlsx"
        result = _convert(workbook, translated_xliff, output)

        assert result.success is False
        assert "bare workbook" in result.errors[0].message
        assert "xliff_map.json" in result.errors[0].message
        assert not output.exists()

    def test_unsupported_map_version_is_rejected(
        self, workbook: Path, tmp_path: Path
    ):
        payload = _map_payload()
        payload["version"] = 99
        skel = tmp_path / "v99.skeleton.zip"
        with zipfile.ZipFile(skel, "w", zipfile.ZIP_DEFLATED) as zf:
            with zipfile.ZipFile(workbook) as source:
                for info in source.infolist():
                    zf.writestr(info, source.read(info))
            zf.writestr(XLIFF_MAP_ENTRY, json.dumps(payload))
        xlf = tmp_path / "t.xlf"
        xlf.write_text(_xliff_text({"1": ("Name", "名称")}), encoding="utf-8")
        output = tmp_path / "out.xlsx"

        result = _convert(skel, xlf, output)
        assert result.success is False
        assert "version 99" in result.errors[0].message
        assert not output.exists()

    def test_map_sheet_absent_from_the_workbook_is_rejected(
        self, workbook: Path, tmp_path: Path
    ):
        payload = _map_payload()
        payload["sheets"] = {"Ghost": payload["sheets"]["Second"]}
        skel = tmp_path / "ghost.skeleton.zip"
        with zipfile.ZipFile(skel, "w", zipfile.ZIP_DEFLATED) as zf:
            with zipfile.ZipFile(workbook) as source:
                for info in source.infolist():
                    zf.writestr(info, source.read(info))
            zf.writestr(XLIFF_MAP_ENTRY, json.dumps(payload))
        xlf = tmp_path / "t.xlf"
        xlf.write_text(_xliff_text({"8": ("Second sheet text", "第二")}), encoding="utf-8")
        output = tmp_path / "out.xlsx"

        result = _convert(skel, xlf, output)
        assert result.success is False
        assert "'Ghost'" in result.errors[0].message
        assert not output.exists()

    def test_invalid_cell_ref_is_rejected(self, workbook: Path, tmp_path: Path):
        payload = _map_payload()
        payload["sheets"]["Second"][0]["cells"][0]["ref"] = "not-a-ref"
        skel = tmp_path / "badref.skeleton.zip"
        with zipfile.ZipFile(skel, "w", zipfile.ZIP_DEFLATED) as zf:
            with zipfile.ZipFile(workbook) as source:
                for info in source.infolist():
                    zf.writestr(info, source.read(info))
            zf.writestr(XLIFF_MAP_ENTRY, json.dumps(payload))
        xlf = tmp_path / "t.xlf"
        xlf.write_text(_xliff_text({"8": ("Second sheet text", "第二")}), encoding="utf-8")
        output = tmp_path / "out.xlsx"

        result = _convert(skel, xlf, output)
        assert result.success is False
        assert "invalid cell ref" in result.errors[0].message
        assert not output.exists()

    def test_missing_xliff_is_reported(self, skeleton: Path, tmp_path: Path):
        output = tmp_path / "out.xlsx"
        result = _convert(skeleton, tmp_path / "absent.xlf", output)

        assert result.success is False
        assert "XLIFF file not found" in result.errors[0].message
        assert not output.exists()


# ── helpers ──────────────────────────────────────────────────────────────

def _source_of(skeleton: Path) -> Path:
    """The workbook the skeleton carries, so tests can read OPP's row text."""
    with zipfile.ZipFile(skeleton) as zf:
        data = {
            info.filename: zf.read(info)
            for info in zf.infolist()
            if info.filename != XLIFF_MAP_ENTRY
        }
    path = skeleton.with_name("extracted.xlsx")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, payload in data.items():
            out.writestr(name, payload)
    return path


def _units_of(xliff: Path) -> dict[str, tuple[str, str | None]]:
    """Read an XLIFF back into {id: (source, target|None)}."""
    from lxml import etree

    units: dict[str, tuple[str, str | None]] = {}
    for element in etree.parse(str(xliff)).getroot().iter():
        if not isinstance(element.tag, str):
            continue
        if etree.QName(element).localname != "trans-unit":
            continue
        source = target = None
        for child in element:
            name = etree.QName(child).localname
            if name == "source":
                source = "".join(child.itertext())
            elif name == "target":
                target = "".join(child.itertext())
        units[element.get("id", "")] = (source or "", target)
    return units